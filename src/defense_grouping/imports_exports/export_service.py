from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from defense_grouping.api.errors import APIError
from defense_grouping.models.availability import Room
from defense_grouping.models.defense import (
    DefenseGroup,
    DefenseSlot,
    PanelAssignment,
    PlanStatus,
    SchedulePlan,
    StudentAssignment,
)
from defense_grouping.models.master import Student, Teacher

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _header(sheet: Worksheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


async def export_plan_workbook(session: AsyncSession, plan: SchedulePlan) -> bytes:
    if plan.status != PlanStatus.PUBLISHED:
        raise APIError(status_code=409, code="plan_not_published", message="只能导出已发布方案")
    groups = list(
        await session.scalars(
            select(DefenseGroup)
            .where(DefenseGroup.plan_id == plan.id)
            .order_by(DefenseGroup.code, DefenseGroup.id)
        )
    )
    group_ids = [group.id for group in groups]
    slots = (
        {
            slot.id: slot
            for slot in await session.scalars(
                select(DefenseSlot).where(DefenseSlot.id.in_(group.slot_id for group in groups))
            )
        }
        if groups
        else {}
    )
    rooms = (
        {
            room.id: room
            for room in await session.scalars(
                select(Room).where(Room.id.in_(group.room_id for group in groups))
            )
        }
        if groups
        else {}
    )
    student_assignments = (
        list(
            await session.scalars(
                select(StudentAssignment).where(StudentAssignment.group_id.in_(group_ids))
            )
        )
        if group_ids
        else []
    )
    panel_assignments = (
        list(
            await session.scalars(
                select(PanelAssignment).where(PanelAssignment.group_id.in_(group_ids))
            )
        )
        if group_ids
        else []
    )
    student_ids = [item.student_id for item in student_assignments]
    teacher_ids = [item.teacher_id for item in panel_assignments]
    students = (
        {
            student.id: student
            for student in await session.scalars(select(Student).where(Student.id.in_(student_ids)))
        }
        if student_ids
        else {}
    )
    teachers = (
        {
            teacher.id: teacher
            for teacher in await session.scalars(select(Teacher).where(Teacher.id.in_(teacher_ids)))
        }
        if teacher_ids
        else {}
    )
    students_by_group: dict[UUID, list[StudentAssignment]] = {}
    panels_by_group: dict[UUID, list[PanelAssignment]] = {}
    for student_assignment in student_assignments:
        students_by_group.setdefault(student_assignment.group_id, []).append(student_assignment)
    for panel_assignment in panel_assignments:
        panels_by_group.setdefault(panel_assignment.group_id, []).append(panel_assignment)

    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "分组总表"
    _header(summary, ["组号", "日期", "开始", "结束", "教室", "组长", "答辩教师", "学生"])
    for group in groups:
        slot = slots[group.slot_id]
        room = rooms[group.room_id]
        panels = panels_by_group.get(group.id, [])
        group_students = students_by_group.get(group.id, [])
        chair = next((item for item in panels if item.is_chair), None)
        summary.append(
            [
                group.code,
                slot.slot_date.isoformat(),
                slot.start_time.isoformat(timespec="minutes"),
                slot.end_time.isoformat(timespec="minutes"),
                f"{room.campus}/{room.building}/{room.name}",
                teachers[chair.teacher_id].name if chair else "",
                "、".join(
                    teachers[item.teacher_id].name
                    for item in sorted(
                        panels, key=lambda value: teachers[value.teacher_id].employee_number
                    )
                ),
                "、".join(
                    f"{students[item.student_id].student_number} {students[item.student_id].name}"
                    for item in sorted(
                        group_students,
                        key=lambda value: students[value.student_id].student_number,
                    )
                ),
            ]
        )

    teacher_sheet = workbook.create_sheet("教师安排")
    _header(teacher_sheet, ["工号", "姓名", "组号", "是否组长", "日期", "时段", "教室"])
    for group in groups:
        slot = slots[group.slot_id]
        room = rooms[group.room_id]
        for panel_assignment in sorted(
            panels_by_group.get(group.id, []),
            key=lambda value: teachers[value.teacher_id].employee_number,
        ):
            teacher = teachers[panel_assignment.teacher_id]
            teacher_sheet.append(
                [
                    teacher.employee_number,
                    teacher.name,
                    group.code,
                    "是" if panel_assignment.is_chair else "否",
                    slot.slot_date.isoformat(),
                    f"{slot.start_time:%H:%M}-{slot.end_time:%H:%M}",
                    f"{room.campus}/{room.building}/{room.name}",
                ]
            )

    student_sheet = workbook.create_sheet("学生安排")
    _header(student_sheet, ["学号", "姓名", "组号", "日期", "时段", "教室"])
    for group in groups:
        slot = slots[group.slot_id]
        room = rooms[group.room_id]
        for student_assignment in sorted(
            students_by_group.get(group.id, []),
            key=lambda value: students[value.student_id].student_number,
        ):
            student = students[student_assignment.student_id]
            student_sheet.append(
                [
                    student.student_number,
                    student.name,
                    group.code,
                    slot.slot_date.isoformat(),
                    f"{slot.start_time:%H:%M}-{slot.end_time:%H:%M}",
                    f"{room.campus}/{room.building}/{room.name}",
                ]
            )

    room_sheet = workbook.create_sheet("教室时段")
    _header(room_sheet, ["日期", "开始", "结束", "校区", "楼宇", "教室", "组号"])
    for group in groups:
        slot = slots[group.slot_id]
        room = rooms[group.room_id]
        room_sheet.append(
            [
                slot.slot_date.isoformat(),
                slot.start_time.isoformat(timespec="minutes"),
                slot.end_time.isoformat(timespec="minutes"),
                room.campus,
                room.building,
                room.name,
                group.code,
            ]
        )

    exception_sheet = workbook.create_sheet("冲突与例外")
    _header(
        exception_sheet,
        ["例外编号", "约束类型", "教师", "学生", "人员", "时段"],
    )
    for item in plan.exception_snapshot:
        exception_sheet.append(
            [
                item.get("id"),
                item.get("constraint_code"),
                item.get("teacher_id"),
                item.get("student_id"),
                item.get("person_id"),
                item.get("slot_id"),
            ]
        )

    metadata = workbook.create_sheet("方案信息")
    _header(metadata, ["字段", "值"])
    metadata.append(["方案版本", plan.version_number])
    metadata.append(["方案编号", str(plan.id)])
    metadata.append(["活动编号", str(plan.activity_id)])
    metadata.append(["状态", plan.status.value])
    metadata.append(["发布时间", plan.published_at.isoformat() if plan.published_at else ""])
    for key, value in sorted(plan.quality_metrics.items()):
        metadata.append([f"质量指标.{key}", value])

    for sheet in workbook.worksheets:
        for column_index, column in enumerate(sheet.columns, start=1):
            width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[get_column_letter(column_index)].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
