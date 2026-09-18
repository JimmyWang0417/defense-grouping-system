from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from defense_grouping.imports_exports.schemas import ImportKind


@dataclass(frozen=True)
class TemplateDefinition:
    filename: str
    sheet_name: str
    headers: tuple[str, ...]
    example: tuple[object, ...]
    descriptions: tuple[str, ...]
    identity_fields: frozenset[str]


TEMPLATES: dict[ImportKind, TemplateDefinition] = {
    ImportKind.STUDENT: TemplateDefinition(
        "学生信息.xlsx",
        "学生信息",
        ("学号", "姓名", "院系", "专业", "年级", "专业方向", "导师工号"),
        ("S2026001", "李明", "软件学院", "软件工程", "2022", "人工智能", "T001"),
        (
            "唯一学号",
            "姓名",
            "院系全称",
            "专业全称",
            "入学年级",
            "多个方向用顿号分隔",
            "导师唯一工号",
        ),
        frozenset({"学号", "导师工号"}),
    ),
    ImportKind.TEACHER: TemplateDefinition(
        "教师信息.xlsx",
        "教师信息",
        ("工号", "姓名", "院系", "职称", "专业方向", "工作量上限"),
        ("T001", "王老师", "软件学院", "副教授", "人工智能", 8),
        ("唯一工号", "姓名", "院系全称", "职称名称", "多个方向用顿号分隔", "最多答辩组数"),
        frozenset({"工号"}),
    ),
    ImportKind.COURSE: TemplateDefinition(
        "课程占用.xlsx",
        "课程占用",
        ("人员编号", "人员类型", "日期或教学周", "星期", "开始节次", "结束节次"),
        ("T001", "教师", "2026-10-12", 1, 1, 2),
        ("工号或学号", "教师/学生", "YYYY-MM-DD 或 学期代码-第N周", "1-7", "起始节次", "结束节次"),
        frozenset({"人员编号"}),
    ),
    ImportKind.LEAVE: TemplateDefinition(
        "请假记录.xlsx",
        "请假记录",
        ("人员编号", "人员类型", "开始时间", "结束时间", "原因"),
        ("T001", "教师", "2026-10-12 08:00", "2026-10-12 12:00", "参加会议"),
        ("工号或学号", "教师/学生", "北京时间", "北京时间", "请假原因"),
        frozenset({"人员编号"}),
    ),
    ImportKind.SLOT: TemplateDefinition(
        "答辩时段.xlsx",
        "答辩时段",
        ("活动", "日期", "开始时间", "结束时间", "最大组数"),
        ("2026 秋季本科答辩", "2026-10-12", "08:00", "10:00", 4),
        ("活动全称", "YYYY-MM-DD", "HH:MM", "HH:MM", "该时段最多并行组数"),
        frozenset({"活动"}),
    ),
    ImportKind.ROOM: TemplateDefinition(
        "教室信息.xlsx",
        "教室信息",
        ("校区", "楼宇", "教室", "容量", "可用日期时段"),
        ("创新港", "泓理楼", "1-101", 30, "2026-10-12 08:00-10:00"),
        ("校区", "楼宇", "教室名称", "人数容量", "YYYY-MM-DD HH:MM-HH:MM；可留空"),
        frozenset({"校区", "楼宇", "教室"}),
    ),
}


@dataclass(frozen=True)
class ParsedRow:
    sheet: str
    row: int
    values: dict[str, Any]


@dataclass(frozen=True)
class ParseResult:
    rows: tuple[ParsedRow, ...]
    structural_errors: tuple[tuple[int, str, str, str], ...]


def generate_template(kind: ImportKind) -> bytes:
    definition = TEMPLATES[kind]
    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.title = definition.sheet_name
    sheet.append(definition.headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(definition.headers))}1"
    for index, header in enumerate(definition.headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(header) * 2 + 4)

    instructions = workbook.create_sheet("填写说明")
    instructions.append(("字段", "说明", "示例"))
    for header, description, example in zip(
        definition.headers,
        definition.descriptions,
        definition.example,
        strict=True,
    ):
        instructions.append((header, description, example))
    instructions.freeze_panes = "A2"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def parse_workbook(path: Path, kind: ImportKind) -> ParseResult:
    definition = TEMPLATES[kind]
    workbook = load_workbook(path, read_only=False, data_only=False, keep_vba=False)
    allowed_sheets = {definition.sheet_name, "填写说明"}
    errors: list[tuple[int, str, str, str]] = []
    unknown_sheets = set(workbook.sheetnames) - allowed_sheets
    if definition.sheet_name not in workbook.sheetnames:
        return ParseResult((), ((1, "", "wrong_sheet", f"缺少工作表：{definition.sheet_name}"),))
    if unknown_sheets:
        errors.append(
            (1, "", "unknown_sheet", f"存在未知工作表：{'、'.join(sorted(unknown_sheets))}")
        )
    sheet = workbook[definition.sheet_name]
    headers = tuple(cell.value for cell in sheet[1])
    if headers != definition.headers:
        return ParseResult(
            (),
            ((1, "", "invalid_headers", f"表头必须为：{'、'.join(definition.headers)}"),),
        )
    if sheet.max_row - 1 > 20_000:
        return ParseResult((), ((1, "", "too_many_rows", "数据行不能超过 20000 行"),))

    rows: list[ParsedRow] = []
    for row_number in range(2, sheet.max_row + 1):
        cells = list(sheet[row_number])
        if all(cell.value in (None, "") for cell in cells):
            continue
        values: dict[str, Any] = {}
        for header, cell in zip(definition.headers, cells, strict=True):
            if header in definition.identity_fields and cell.data_type == "f":
                errors.append(
                    (row_number, header, "formula_not_allowed", "业务编号字段不能使用公式")
                )
            value = cell.value
            values[header] = value.strip() if isinstance(value, str) else value
        rows.append(ParsedRow(definition.sheet_name, row_number, values))
    return ParseResult(tuple(rows), tuple(errors))
