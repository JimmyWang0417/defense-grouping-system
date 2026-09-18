import asyncio
from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from defense_grouping.models.governance import AuditLog, ImportBatch

STUDENT_HEADERS = ["学号", "姓名", "院系", "专业", "年级", "专业方向", "导师工号"]


def workbook_bytes(sheet_name: str, headers: list[str], rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


async def wait_for_import(master_harness, batch_id: str) -> dict[str, object]:
    headers = await master_harness.headers()
    for _attempt in range(100):
        response = await master_harness.client.get(
            f"/api/v1/imports/{batch_id}",
            headers=headers,
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        if payload["status"] in {"ready", "failed", "confirmed"}:
            return payload
        await asyncio.sleep(0.02)
    raise AssertionError("import preview did not reach a terminal state")


@pytest.mark.asyncio
async def test_six_templates_have_locked_headers(master_harness) -> None:
    expected = {
        "student": ("学生信息", STUDENT_HEADERS),
        "teacher": ("教师信息", ["工号", "姓名", "院系", "职称", "专业方向", "工作量上限"]),
        "course": ("课程占用", ["人员编号", "人员类型", "日期或教学周", "星期", "开始节次", "结束节次"]),
        "leave": ("请假记录", ["人员编号", "人员类型", "开始时间", "结束时间", "原因"]),
        "slot": ("答辩时段", ["活动", "日期", "开始时间", "结束时间", "最大组数"]),
        "room": ("教室信息", ["校区", "楼宇", "教室", "容量", "可用日期时段"]),
    }
    headers = await master_harness.headers()

    for kind, (sheet_name, columns) in expected.items():
        response = await master_harness.client.get(
            f"/api/v1/imports/templates/{kind}",
            headers=headers,
        )
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        assert workbook.sheetnames == [sheet_name, "填写说明"]
        assert [cell.value for cell in next(workbook[sheet_name].iter_rows())] == columns


@pytest.mark.asyncio
async def test_student_preflight_reports_exact_cell(master_harness) -> None:
    headers = await master_harness.headers()
    await prepare_student_dependencies(master_harness, headers)
    content = workbook_bytes(
        "学生信息",
        STUDENT_HEADERS,
        [
            [None, None, None, None, None, None, None],
            ["S001", "李明", "软件学院", "软件工程", "2022", "人工智能", "UNKNOWN"],
        ],
    )
    created = await master_harness.client.post(
        "/api/v1/imports/student/preflight",
        headers=headers,
        files={
            "file": (
                "学生信息.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert created.status_code == 202

    preview = await wait_for_import(master_harness, created.json()["id"])
    issue = preview["issues"][0]
    assert issue["sheet"] == "学生信息"
    assert issue["row"] == 3
    assert issue["field"] == "导师工号"
    assert issue["severity"] == "error"


async def prepare_student_dependencies(master_harness, headers):
    major = await master_harness.client.post(
        "/api/v1/majors",
        headers=headers,
        json={
            "department_id": str(master_harness.department_id),
            "code": "SE",
            "name": "软件工程",
        },
    )
    direction = await master_harness.client.post(
        "/api/v1/directions",
        headers=headers,
        json={
            "department_id": str(master_harness.department_id),
            "code": "AI",
            "name": "人工智能",
        },
    )
    teacher = await master_harness.client.post(
        "/api/v1/teachers",
        headers=headers,
        json={
            "employee_number": "T001",
            "name": "导师甲",
            "department_id": str(master_harness.department_id),
            "title": "副教授",
            "title_rank": 4,
            "direction_ids": [direction.json()["id"]],
            "workload_limit": 10,
        },
    )
    assert major.status_code == direction.status_code == teacher.status_code == 201


async def create_valid_student_preview(master_harness, headers):
    content = workbook_bytes(
        "学生信息",
        STUDENT_HEADERS,
        [["S001", "李明", "软件学院", "软件工程", "2022", "人工智能", "T001"]],
    )
    created = await master_harness.client.post(
        "/api/v1/imports/student/preflight",
        headers=headers,
        files={"file": ("学生信息.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert created.status_code == 202
    preview = await wait_for_import(master_harness, created.json()["id"])
    assert preview["status"] == "ready"
    assert preview["issues"] == []
    return preview, content


async def preflight_and_confirm(master_harness, headers, kind, filename, content):
    created = await master_harness.client.post(
        f"/api/v1/imports/{kind}/preflight",
        headers=headers,
        files={
            "file": (
                filename,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert created.status_code == 202, created.text
    preview = await wait_for_import(master_harness, created.json()["id"])
    assert preview["status"] == "ready", preview
    confirmed = await master_harness.client.post(
        f"/api/v1/imports/{preview['id']}/confirm",
        headers=headers,
    )
    assert confirmed.status_code == 200, confirmed.text
    return confirmed.json()


@pytest.mark.asyncio
async def test_confirm_is_transactional_and_repeated_import_is_idempotent(master_harness) -> None:
    headers = await master_harness.headers()
    await prepare_student_dependencies(master_harness, headers)
    preview, content = await create_valid_student_preview(master_harness, headers)

    confirmed = await master_harness.client.post(
        f"/api/v1/imports/{preview['id']}/confirm",
        headers=headers,
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["creates"] == 1
    students = await master_harness.client.get("/api/v1/students", headers=headers)
    assert students.json()["total"] == 1
    async with master_harness.database.session() as session:
        audit = await session.scalar(
            select(AuditLog).where(
                AuditLog.action == "import.confirmed",
                AuditLog.object_id == UUID(str(preview["id"])),
            )
        )
        assert audit is not None

    repeated = await master_harness.client.post(
        "/api/v1/imports/student/preflight",
        headers=headers,
        files={"file": ("学生信息.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    second = await wait_for_import(master_harness, repeated.json()["id"])
    assert second["updates"] == 1
    reconfirmed = await master_harness.client.post(
        f"/api/v1/imports/{second['id']}/confirm",
        headers=headers,
    )
    assert reconfirmed.status_code == 200
    students = await master_harness.client.get("/api/v1/students", headers=headers)
    assert students.json()["total"] == 1


@pytest.mark.asyncio
async def test_teacher_course_leave_slot_and_room_imports(master_harness) -> None:
    headers = await master_harness.headers()
    await prepare_student_dependencies(master_harness, headers)

    teacher = workbook_bytes(
        "教师信息",
        ["工号", "姓名", "院系", "职称", "专业方向", "工作量上限"],
        [["T002", "教师乙", "软件学院", "讲师", "人工智能", 6]],
    )
    await preflight_and_confirm(master_harness, headers, "teacher", "教师信息.xlsx", teacher)

    course = workbook_bytes(
        "课程占用",
        ["人员编号", "人员类型", "日期或教学周", "星期", "开始节次", "结束节次"],
        [["T002", "教师", "2026-10-12", 1, 1, 2]],
    )
    await preflight_and_confirm(master_harness, headers, "course", "课程占用.xlsx", course)

    leave = workbook_bytes(
        "请假记录",
        ["人员编号", "人员类型", "开始时间", "结束时间", "原因"],
        [["T002", "教师", "2026-10-12 13:00", "2026-10-12 14:00", "参加会议"]],
    )
    await preflight_and_confirm(master_harness, headers, "leave", "请假记录.xlsx", leave)

    activity = await master_harness.client.post(
        "/api/v1/activities",
        headers=headers,
        json={
            "department_id": str(master_harness.department_id),
            "name": "导入测试答辩",
            "academic_year": "2026-2027",
            "semester": "秋",
        },
    )
    assert activity.status_code == 201
    slot = workbook_bytes(
        "答辩时段",
        ["活动", "日期", "开始时间", "结束时间", "最大组数"],
        [["导入测试答辩", "2026-10-12", "08:00", "10:00", 4]],
    )
    await preflight_and_confirm(master_harness, headers, "slot", "答辩时段.xlsx", slot)

    room = workbook_bytes(
        "教室信息",
        ["校区", "楼宇", "教室", "容量", "可用日期时段"],
        [["创新港", "泓理楼", "1-101", 30, "2026-10-12 08:00-10:00"]],
    )
    await preflight_and_confirm(master_harness, headers, "room", "教室信息.xlsx", room)


@pytest.mark.asyncio
async def test_preflight_rejects_duplicate_rows_wrong_department_and_macros(master_harness) -> None:
    headers = await master_harness.headers()
    await prepare_student_dependencies(master_harness, headers)
    duplicate = workbook_bytes(
        "学生信息",
        STUDENT_HEADERS,
        [
            ["S001", "李明", "软件学院", "软件工程", "2022", "人工智能", "T001"],
            ["S001", "李明", "软件学院", "软件工程", "2022", "人工智能", "T001"],
        ],
    )
    created = await master_harness.client.post(
        "/api/v1/imports/student/preflight",
        headers=headers,
        files={"file": ("学生信息.xlsx", duplicate, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    preview = await wait_for_import(master_harness, created.json()["id"])
    assert any(item["code"] == "duplicate_row" for item in preview["issues"])

    wrong_department = workbook_bytes(
        "教师信息",
        ["工号", "姓名", "院系", "职称", "专业方向", "工作量上限"],
        [["T009", "越权教师", "计算机学院", "讲师", "人工智能", 6]],
    )
    created = await master_harness.client.post(
        "/api/v1/imports/teacher/preflight",
        headers=headers,
        files={"file": ("教师信息.xlsx", wrong_department, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    preview = await wait_for_import(master_harness, created.json()["id"])
    assert any(item["code"] == "department_mismatch" for item in preview["issues"])

    macro = await master_harness.client.post(
        "/api/v1/imports/student/preflight",
        headers=headers,
        files={"file": ("学生信息.xlsm", duplicate, "application/vnd.ms-excel.sheet.macroEnabled.12")},
    )
    assert macro.status_code == 422
    assert macro.json()["error"]["code"] == "invalid_workbook_type"

@pytest.mark.asyncio
async def test_confirm_rejects_file_changed_after_preview(master_harness) -> None:
    headers = await master_harness.headers()
    await prepare_student_dependencies(master_harness, headers)
    preview, _content = await create_valid_student_preview(master_harness, headers)
    async with master_harness.database.session() as session:
        batch = await session.get(ImportBatch, UUID(str(preview["id"])))
        assert batch is not None
        Path(batch.stored_path).write_bytes(b"changed")

    response = await master_harness.client.post(
        f"/api/v1/imports/{preview['id']}/confirm",
        headers=headers,
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "import_file_changed"


@pytest.mark.asyncio
async def test_confirm_rolls_back_every_row_when_write_fails(master_harness) -> None:
    from defense_grouping.imports_exports.routes import get_import_writer

    class FailingImportWriter:
        async def apply(self, *args, **kwargs):
            del args, kwargs
            raise RuntimeError("injected row failure")

    headers = await master_harness.headers()
    await prepare_student_dependencies(master_harness, headers)
    preview, _content = await create_valid_student_preview(master_harness, headers)
    master_harness.app.dependency_overrides[get_import_writer] = lambda: FailingImportWriter()
    try:
        response = await master_harness.client.post(
            f"/api/v1/imports/{preview['id']}/confirm",
            headers=headers,
        )
    finally:
        master_harness.app.dependency_overrides.pop(get_import_writer, None)

    assert response.status_code == 500
    students = await master_harness.client.get("/api/v1/students", headers=headers)
    assert students.json()["total"] == 0
