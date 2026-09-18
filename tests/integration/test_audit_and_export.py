from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from defense_grouping.governance.audit import (
    AuditMutationError,
    redact_summary,
    write_audit,
)
from defense_grouping.models.governance import AuditLog
from tests.integration.schedule_support import create_ready_plan


@pytest.mark.asyncio
async def test_export_matches_published_plan(master_harness, configured_schedule) -> None:
    headers, ready = await create_ready_plan(master_harness, configured_schedule)
    published = await master_harness.client.post(
        f"/api/v1/plans/{ready['id']}/publish",
        headers=headers,
    )
    assert published.status_code == 200

    response = await master_harness.client.get(
        f"/api/v1/plans/{ready['id']}/export",
        headers=headers,
    )

    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["分组总表", "教师安排", "学生安排", "教室时段", "冲突与例外", "方案信息"]
    assert workbook["分组总表"]["A2"].value == published.json()["groups"][0]["code"]
    assert workbook["方案信息"]["B2"].value == published.json()["version_number"]


@pytest.mark.asyncio
async def test_audit_redaction_and_query(master_harness) -> None:
    summary = {
        "username": "admin",
        "password": "secret",
        "nested": {"token": "jwt", "safe": "visible"},
        "items": [{"refresh_token": "refresh"}],
    }
    assert redact_summary(summary) == {
        "username": "admin",
        "password": "[REDACTED]",
        "nested": {"token": "[REDACTED]", "safe": "visible"},
        "items": [{"refresh_token": "[REDACTED]"}],
    }
    async with master_harness.database.session() as session:
        await write_audit(
            session,
            actor_id=None,
            action="test.redaction",
            object_type="test",
            object_id=None,
            request_id="audit-redaction-request",
            before_summary=None,
            after_summary=summary,
        )
        await session.commit()

    response = await master_harness.client.get(
        "/api/v1/audit-logs?action=test.redaction",
        headers=await master_harness.headers("system-admin"),
    )

    assert response.status_code == 200, response.text
    assert response.json()["items"][0]["after_summary"]["password"] == "[REDACTED]"
    assert response.json()["items"][0]["after_summary"]["nested"]["token"] == "[REDACTED]"


@pytest.mark.asyncio
async def test_successful_write_request_has_request_id_audit(master_harness) -> None:
    request_id = "master-write-audit"
    response = await master_harness.client.post(
        "/api/v1/departments",
        headers={**(await master_harness.headers("system-admin")), "X-Request-ID": request_id},
        json={"code": "AUDIT", "name": "审计测试学院"},
    )
    assert response.status_code == 201

    logs = await master_harness.client.get(
        f"/api/v1/audit-logs?request_id={request_id}",
        headers=await master_harness.headers("system-admin"),
    )
    assert logs.status_code == 200
    assert logs.json()["total"] == 1
    assert logs.json()["items"][0]["request_id"] == request_id


@pytest.mark.asyncio
async def test_audit_records_cannot_be_updated_or_deleted(master_harness) -> None:
    async with master_harness.database.session() as session:
        record = await write_audit(
            session,
            actor_id=None,
            action="test.append-only",
            object_type="test",
            object_id=None,
            request_id="append-only-request",
            before_summary=None,
            after_summary={"value": "original"},
        )
        await session.commit()
        record_id = record.id

    async with master_harness.database.session() as session:
        record = await session.scalar(select(AuditLog).where(AuditLog.id == record_id))
        assert record is not None
        record.action = "test.mutated"
        with pytest.raises(AuditMutationError, match="cannot be updated"):
            await session.flush()
        await session.rollback()

    async with master_harness.database.session() as session:
        record = await session.scalar(select(AuditLog).where(AuditLog.id == record_id))
        assert record is not None
        await session.delete(record)
        with pytest.raises(AuditMutationError, match="cannot be deleted"):
            await session.flush()
