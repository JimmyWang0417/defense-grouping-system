from __future__ import annotations

import asyncio
import os
import subprocess
import sys
from io import BytesIO

import httpx2
import pytest
from openpyxl import Workbook, load_workbook
from testcontainers.community.postgres import PostgresContainer

from defense_grouping.api.main import create_app
from defense_grouping.auth.security import hash_password
from defense_grouping.config import Settings
from defense_grouping.db.session import Database
from defense_grouping.models.identity import Role, RoleAssignment, User

PASSWORD = "Postgres Admin Password 2026"

pytestmark = [
    pytest.mark.postgres,
    pytest.mark.skipif(
        os.getenv("RUN_POSTGRES_TESTS") != "1",
        reason="set RUN_POSTGRES_TESTS=1 on a host with Docker",
    ),
]


def _student_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生信息"
    sheet.append(["学号", "姓名", "院系", "专业", "年级", "专业方向", "导师工号"])
    sheet.append(["PG-S001", "数据库学生", "软件学院", "软件工程", "2022", "智能软件", "PG-T001"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


async def _wait_for_import(
    client: httpx2.AsyncClient, headers: dict[str, str], batch_id: str
) -> dict[str, object]:
    for _attempt in range(200):
        response = await client.get(f"/api/v1/imports/{batch_id}", headers=headers)
        assert response.status_code == 200, response.text
        body = response.json()
        if body["status"] in {"ready", "failed"}:
            return body
        await asyncio.sleep(0.05)
    raise AssertionError("PostgreSQL import preflight timed out")


async def _wait_for_schedule(
    client: httpx2.AsyncClient, headers: dict[str, str], job_id: str
) -> dict[str, object]:
    for _attempt in range(300):
        response = await client.get(f"/api/v1/schedule-jobs/{job_id}", headers=headers)
        assert response.status_code == 200, response.text
        body = response.json()
        if body["status"] in {"succeeded", "failed", "cancelled"}:
            return body
        await asyncio.sleep(0.05)
    raise AssertionError("PostgreSQL schedule job timed out")


@pytest.mark.asyncio
async def test_postgres_migrations_and_core_workflow(tmp_path) -> None:
    with PostgresContainer("postgres:17-alpine", driver="asyncpg") as postgres:
        database_url = postgres.get_connection_url(driver="asyncpg")
        migration_environment = {**os.environ, "DEFENSE_DATABASE_URL": database_url}
        migration = await asyncio.to_thread(
            subprocess.run,
            [sys.executable, "-m", "alembic", "upgrade", "head"],
            check=False,
            capture_output=True,
            cwd=os.getcwd(),
            env=migration_environment,
            text=True,
        )
        assert migration.returncode == 0, migration.stderr

        seed_database = Database(database_url)
        async with seed_database.session() as session:
            admin = User(
                username="postgres-admin",
                password_hash=hash_password(PASSWORD),
                is_active=True,
                must_change_password=False,
            )
            session.add(admin)
            await session.flush()
            session.add(
                RoleAssignment(
                    user_id=admin.id,
                    role=Role.SYSTEM_ADMIN,
                    department_id=None,
                    can_approve_exceptions=False,
                )
            )
            await session.commit()
        await seed_database.dispose()

        settings = Settings(
            environment="test",
            database_url=database_url,
            jwt_secret="postgres-test-secret-with-at-least-32-characters",
            storage_dir=tmp_path / "storage",
        )
        app = create_app(settings)
        transport = httpx2.ASGITransport(app=app)
        async with httpx2.AsyncClient(transport=transport, base_url="http://testserver") as client:
            login = await client.post(
                "/api/v1/auth/login",
                json={"username": "postgres-admin", "password": PASSWORD},
            )
            assert login.status_code == 200, login.text
            headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

            department_response = await client.post(
                "/api/v1/departments",
                headers=headers,
                json={"code": "PG-SE", "name": "软件学院"},
            )
            assert department_response.status_code == 201, department_response.text
            department_id = department_response.json()["id"]

            major = await client.post(
                "/api/v1/majors",
                headers=headers,
                json={
                    "department_id": department_id,
                    "code": "PG-SE",
                    "name": "软件工程",
                },
            )
            direction = await client.post(
                "/api/v1/directions",
                headers=headers,
                json={
                    "department_id": department_id,
                    "code": "PG-AI",
                    "name": "智能软件",
                },
            )
            assert major.status_code == direction.status_code == 201
            direction_id = direction.json()["id"]

            for number, name, rank in (
                ("PG-T001", "数据库导师", 5),
                ("PG-T002", "数据库委员甲", 5),
                ("PG-T003", "数据库委员乙", 4),
            ):
                teacher = await client.post(
                    "/api/v1/teachers",
                    headers=headers,
                    json={
                        "employee_number": number,
                        "name": name,
                        "department_id": department_id,
                        "title": "教授" if rank == 5 else "副教授",
                        "title_rank": rank,
                        "direction_ids": [direction_id],
                        "workload_limit": 5,
                    },
                )
                assert teacher.status_code == 201, teacher.text

            preview = await client.post(
                f"/api/v1/imports/student/preflight?department_id={department_id}",
                headers=headers,
                files={
                    "file": (
                        "学生信息.xlsx",
                        _student_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert preview.status_code == 202, preview.text
            batch = await _wait_for_import(client, headers, preview.json()["id"])
            assert batch["status"] == "ready", batch
            confirmed = await client.post(
                f"/api/v1/imports/{preview.json()['id']}/confirm", headers=headers
            )
            assert confirmed.status_code == 200, confirmed.text
            assert confirmed.json()["creates"] == 1

            activity = await client.post(
                "/api/v1/activities",
                headers=headers,
                json={
                    "department_id": department_id,
                    "name": "PostgreSQL 兼容性答辩",
                    "academic_year": "2026-2027",
                    "semester": "秋",
                },
            )
            room = await client.post(
                "/api/v1/rooms",
                headers=headers,
                json={
                    "department_id": department_id,
                    "campus": "创新港",
                    "building": "泓理楼",
                    "name": "PG-101",
                    "capacity": 20,
                },
            )
            assert activity.status_code == room.status_code == 201
            activity_id = activity.json()["id"]
            slot = await client.post(
                f"/api/v1/activities/{activity_id}/slots",
                headers=headers,
                json={
                    "date": "2026-10-12",
                    "start_time": "08:00",
                    "end_time": "10:00",
                    "max_groups": 1,
                },
            )
            rules = await client.put(
                f"/api/v1/activities/{activity_id}/rules",
                headers=headers,
                json={
                    "students_per_group": 1,
                    "teachers_per_group": 2,
                    "chair_min_title_rank": 4,
                    "teacher_workload_limit": 5,
                    "balance_students_weight": 100,
                    "balance_teacher_load_weight": 50,
                    "direction_match_weight": 30,
                    "compact_schedule_weight": 20,
                    "exception_use_weight": 500,
                },
            )
            rooms = await client.put(
                f"/api/v1/activities/{activity_id}/rooms",
                headers=headers,
                json={"room_ids": [room.json()["id"]]},
            )
            assert slot.status_code == 201, slot.text
            assert rules.status_code == rooms.status_code == 200

            started = await client.post(
                f"/api/v1/activities/{activity_id}/schedule-jobs",
                headers={**headers, "Idempotency-Key": "postgres-core-flow"},
                json={"seed": 20260918, "time_limit_seconds": 5},
            )
            assert started.status_code == 202, started.text
            job = await _wait_for_schedule(client, headers, started.json()["id"])
            assert job["status"] == "succeeded", job
            plan_id = str(job["plan_id"])

            published = await client.post(f"/api/v1/plans/{plan_id}/publish", headers=headers)
            assert published.status_code == 200, published.text
            assert published.json()["status"] == "published"
            exported = await client.get(f"/api/v1/plans/{plan_id}/export", headers=headers)
            assert exported.status_code == 200, exported.text
            workbook = load_workbook(BytesIO(exported.content), read_only=True)
            assert "分组总表" in workbook.sheetnames

        await app.state.task_executor.shutdown()
        await app.state.database.dispose()
